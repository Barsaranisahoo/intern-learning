import csv
import io

import markdown
import openpyxl
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from odf import teletype
from odf.opendocument import load
from pptx import Presentation
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text


def extract_text(file, filename):

    filename = filename.lower()

    # ---------------- TXT ----------------

    if filename.endswith(".txt"):

        content = file.read()

        return content.decode("utf-8", errors="ignore")

    # ---------------- PDF ----------------

    elif filename.endswith(".pdf"):

        reader = PdfReader(file)

        text = ""

        for page in reader.pages:

            text += (page.extract_text() or "") + "\n"

        return text

    # ---------------- DOCX ----------------

    elif filename.endswith(".docx"):

        doc = Document(file)

        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    # ---------------- PPTX ----------------

    elif filename.endswith(".pptx"):

        presentation = Presentation(file)

        text = ""

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text += shape.text + "\n"

        return text

    # ---------------- CSV ----------------

    elif filename.endswith(".csv"):

        file.seek(0)

        decoded = io.StringIO(
            file.read().decode("utf-8", errors="ignore")
        )

        reader = csv.reader(decoded)

        rows = []

        for row in reader:

            rows.append(" | ".join(row))

        return "\n".join(rows)

    # ---------------- XLSX ----------------

    elif filename.endswith(".xlsx"):

        workbook = openpyxl.load_workbook(
            file,
            data_only=True
        )

        text = ""

        for sheet in workbook.worksheets:

            text += f"\nSheet: {sheet.title}\n"

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:

                    text += " | ".join(values) + "\n"

        return text

    # ---------------- XLS ----------------

    elif filename.endswith(".xls"):

        data = pd.read_excel(file)

        return data.to_string(index=False)

    # ---------------- HTML ----------------

    elif filename.endswith(".html") or filename.endswith(".htm"):

        html = file.read().decode(
            "utf-8",
            errors="ignore"
        )

        soup = BeautifulSoup(html, "html.parser")

        return soup.get_text(separator="\n")

    # ---------------- Markdown ----------------

    elif filename.endswith(".md"):

        md = file.read().decode(
            "utf-8",
            errors="ignore"
        )

        html = markdown.markdown(md)

        soup = BeautifulSoup(html, "html.parser")

        return soup.get_text(separator="\n")

    # ---------------- RTF ----------------

    elif filename.endswith(".rtf"):

        rtf = file.read().decode(
            "utf-8",
            errors="ignore"
        )

        return rtf_to_text(rtf)

    # ---------------- ODT ----------------

    elif filename.endswith(".odt"):

        document = load(file)

        return teletype.extractText(document)

    # ---------------- ODS ----------------

    elif filename.endswith(".ods"):

        data = pd.read_excel(
            file,
            engine="odf"
        )

        return data.to_string(index=False)
    # ---------------- SOURCE CODE ----------------

        # ---------------- SOURCE CODE ----------------

    elif filename.endswith(
        (
            ".py",
            ".js",
            ".ts",
            ".tsx",
            ".jsx",
            ".java",
            ".cpp",
            ".c",
            ".h",
            ".hpp",
            ".cs",
            ".go",
            ".php",
            ".rb",
            ".swift",
            ".kt",
            ".rs",
            ".css",
            ".scss",
            ".sass",
            ".json",
            ".xml",
            ".sql",
            ".yaml",
            ".yml",
            ".sh",
            ".bat",
            ".ps1",
        )
    ) or filename in (
        "dockerfile",
        "requirements.txt",
        "package.json",
        "package-lock.json",
        "pyproject.toml",
        "vite.config.ts",
        "tsconfig.json",
        ".env.example",
    ):

        file.seek(0)

        try:
            return file.read().decode("utf-8")
        except UnicodeDecodeError:
            file.seek(0)
            return file.read().decode(
                "latin-1",
                errors="ignore",
            )
    else:

        raise ValueError(
            f"Unsupported file type: {filename}"
        )